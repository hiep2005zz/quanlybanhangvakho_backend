import io
import pytest
import openpyxl
from fastapi.testclient import TestClient
from app.main import app

client = TestClient(app)


def get_token(username: str = "admin", password: str = "123") -> str:
    resp = client.post("/api/v1/auth/login", json={"username": username, "password": password})
    assert resp.status_code == 200, f"Login failed: {resp.text}"
    return resp.json()["access_token"]


def test_turnover_report_unauthorized():
    """Kiểm tra chặn truy cập khi chưa đăng nhập."""
    resp = client.get("/api/v1/inventory/turnover-report")
    assert resp.status_code in [401, 403]

    resp_exp = client.get("/api/v1/inventory/turnover-report/export")
    assert resp_exp.status_code in [401, 403]


def test_turnover_report_structure_and_kpis():
    """Kiểm tra cấu trúc phản hồi báo cáo và các thẻ KPI tổng hợp."""
    token = get_token("admin")
    headers = {"Authorization": f"Bearer {token}"}

    resp = client.get("/api/v1/inventory/turnover-report", headers=headers)
    assert resp.status_code == 200
    data = resp.json()

    assert "summary" in data
    assert "warehouse_breakdown" in data
    assert "category_breakdown" in data
    assert "items" in data
    assert "warehouses" in data
    assert "categories" in data
    assert "total" in data

    # 1. Kiểm tra KPI Summary
    summary = data["summary"]
    assert "total_inventory_value" in summary
    assert "slow_stagnant_value" in summary
    assert "slow_stagnant_skus" in summary
    assert "total_skus" in summary
    assert "avg_turnover_rate" in summary
    assert "avg_dsi" in summary
    assert "total_cogs" in summary
    assert summary["period_days"] == 30

    assert summary["total_inventory_value"] >= 0
    assert summary["slow_stagnant_value"] >= 0
    assert summary["slow_stagnant_skus"] <= summary["total_skus"]

    # 2. Kiểm tra phân bổ kho & ngành hàng
    assert len(data["warehouse_breakdown"]) > 0
    for wh in data["warehouse_breakdown"]:
        assert "warehouse_id" in wh
        assert "warehouse_name" in wh
        assert "total_value" in wh
        assert "percentage" in wh

    assert len(data["category_breakdown"]) > 0
    for cat in data["category_breakdown"]:
        assert "category" in cat
        assert "total_value" in cat
        assert "percentage" in cat

    # 3. Kiểm tra từng dòng SKU
    assert len(data["items"]) > 0
    for item in data["items"]:
        assert "product_id" in item
        assert "product_code" in item
        assert "product_name" in item
        assert "category" in item
        assert "base_unit" in item
        assert "current_stock" in item
        assert "cost_price" in item
        assert "inventory_value" in item
        assert "sold_quantity" in item
        assert "cogs" in item
        assert "turnover_rate" in item
        assert "status" in item
        assert "status_label" in item
        assert "status_severity" in item


def test_turnover_report_zero_division_and_formulas():
    """
    Kiểm tra xử lý ngoại lệ Zero-Division và công thức tính toán:
    - Nếu COGS = 0 hoặc hàng không bán được gì -> DSI phải trả về null/None (không crash).
    - Status phải là 'stagnant' hoặc 'out_of_stock'.
    - COGS = sold_quantity * cost_price.
    """
    token = get_token("admin")
    headers = {"Authorization": f"Bearer {token}"}

    resp = client.get("/api/v1/inventory/turnover-report?days=30", headers=headers)
    assert resp.status_code == 200
    data = resp.json()

    for item in data["items"]:
        # Kiểm tra COGS = sold_quantity * cost_price
        expected_cogs = round(item["sold_quantity"] * item["cost_price"], 2)
        assert abs(item["cogs"] - expected_cogs) < 0.1

        # Kiểm tra Zero-Division
        if item["sold_quantity"] == 0 and item["current_stock"] > 0:
            assert item["turnover_rate"] == 0.0
            assert item["dsi"] is None
            assert item["status"] == "stagnant"
            assert "Tồn bất động" in item["status_label"]
            assert item["status_severity"] == "danger"
        elif item["current_stock"] > 0 and item["sold_quantity"] > 0:
            # Có bán và có tồn -> DSI phải là số dương
            if item["turnover_rate"] > 0:
                assert item["dsi"] is not None
                assert item["dsi"] > 0


def test_turnover_report_filters():
    """Kiểm tra các bộ lọc: kho, ngành hàng, chu kỳ số ngày, tìm kiếm và trạng thái."""
    token = get_token("admin")
    headers = {"Authorization": f"Bearer {token}"}

    # 1. Lọc theo kho WH01
    resp_wh = client.get("/api/v1/inventory/turnover-report?warehouse_id=WH01", headers=headers)
    assert resp_wh.status_code == 200
    data_wh = resp_wh.json()
    for item in data_wh["items"]:
        assert item["warehouse_id"] == "WH01"

    # 2. Lọc theo chu kỳ 60 ngày
    resp_60 = client.get("/api/v1/inventory/turnover-report?days=60", headers=headers)
    assert resp_60.status_code == 200
    data_60 = resp_60.json()
    assert data_60["summary"]["period_days"] == 60

    # 3. Lọc theo trạng thái ứ đọng
    resp_stagnant = client.get("/api/v1/inventory/turnover-report?status=stagnant", headers=headers)
    assert resp_stagnant.status_code == 200
    data_stagnant = resp_stagnant.json()
    for item in data_stagnant["items"]:
        assert item["status"] == "stagnant"

    # 4. Tìm kiếm theo SKU
    resp_search = client.get("/api/v1/inventory/turnover-report?search=SP001", headers=headers)
    assert resp_search.status_code == 200
    data_search = resp_search.json()
    for item in data_search["items"]:
        assert "SP001" in item["product_code"] or "SP001" in item["product_name"]


def test_turnover_report_excel_export():
    """Kiểm tra xuất file Excel (.xlsx) chuẩn nghiệp vụ 2 Sheet."""
    token = get_token("admin")
    headers = {"Authorization": f"Bearer {token}"}

    resp = client.get("/api/v1/inventory/turnover-report/export?days=30", headers=headers)
    assert resp.status_code == 200
    assert "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet" in resp.headers["content-type"]
    assert "attachment; filename=" in resp.headers["content-disposition"]
    assert resp.headers["content-disposition"].endswith('.xlsx"')

    # Đọc lại nội dung file Excel bằng openpyxl để xác thực cấu trúc
    excel_stream = io.BytesIO(resp.content)
    wb = openpyxl.load_workbook(excel_stream)
    assert "Chi Tiết SKU & Vòng Quay" in wb.sheetnames
    assert "Phân Bổ Tồn Kho" in wb.sheetnames

    ws1 = wb["Chi Tiết SKU & Vòng Quay"]
    assert "BÁO CÁO TỒN KHO & VÒNG QUAY HÀNG HÓA" in str(ws1["A1"].value)
    # Header row
    assert ws1["A9"].value == "STT"
    assert ws1["B9"].value == "Mã SKU"
    assert ws1["C9"].value == "Tên sản phẩm"
    assert ws1["H9"].value == "Giá trị tồn (đ)"

    ws2 = wb["Phân Bổ Tồn Kho"]
    assert "BÁO CÁO PHÂN BỔ GIÁ TRỊ TỒN KHO" in str(ws2["A1"].value)


def test_turnover_report_excel_export_by_warehouse():
    """Kiểm tra xuất file Excel lọc theo từng kho cụ thể (KHO_HN) và toàn bộ kho."""
    token = get_token("admin")
    headers = {"Authorization": f"Bearer {token}"}

    # 1. Xuất theo kho KHO_HN
    resp_hn = client.get("/api/v1/inventory/turnover-report/export?warehouse_id=KHO_HN&days=30", headers=headers)
    assert resp_hn.status_code == 200
    assert "Bao_Cao_Ton_Kho_Vong_Quay_KhoTongHaNoi_" in resp_hn.headers["content-disposition"]

    wb_hn = openpyxl.load_workbook(io.BytesIO(resp_hn.content))
    ws1_hn = wb_hn["Chi Tiết SKU & Vòng Quay"]
    assert ws1_hn["B4"].value == "Kho Tổng Hà Nội"

    ws2_hn = wb_hn["Phân Bổ Tồn Kho"]
    wh_hn_found = False
    for r in range(5, 10):
        val = ws2_hn.cell(row=r, column=1).value
        if val == "KHO_HN":
            assert "(Đang chọn)" in ws2_hn.cell(row=r, column=2).value
            assert ws2_hn.cell(row=r, column=1).fill.start_color.rgb == "00DCFCE7"
            wh_hn_found = True
            break
    assert wh_hn_found

    # 2. Xuất toàn bộ kho
    resp_all = client.get("/api/v1/inventory/turnover-report/export?warehouse_id=all&days=30", headers=headers)
    assert resp_all.status_code == 200
    assert "Bao_Cao_Ton_Kho_Vong_Quay_Tat_Ca_" in resp_all.headers["content-disposition"]
    wb_all = openpyxl.load_workbook(io.BytesIO(resp_all.content))
    ws1_all = wb_all["Chi Tiết SKU & Vòng Quay"]
    assert ws1_all["B4"].value == "Tất cả các kho"
