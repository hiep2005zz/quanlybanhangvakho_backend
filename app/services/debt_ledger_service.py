import io
from datetime import datetime, timezone
from typing import List, Optional, Tuple
from sqlalchemy.orm import Session
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

from app.models.entities import DealerEntity, OrderEntity, UserEntity
from app.models.invoice import InvoiceEntity
from app.schemas.debt_ledger import (
    AgeingBuckets,
    DealerDebtAgeingItem,
    DebtLedgerSummary,
    DebtLedgerResponse,
)

def compute_debt_ledger_data(
    db: Session,
    region: Optional[str] = None,
    assigned_sale_id: Optional[int] = None,
    overdue_only: bool = False,
    search: Optional[str] = None,
    ageing_bucket: Optional[str] = None,
) -> DebtLedgerResponse:
    """
    Computes debt ledger per dealer with 4 ageing buckets:
    - 0-30 days
    - 31-60 days
    - 61-90 days
    - > 90 days
    Directly calculated from real database orders, confirmed payment collections, and dealers.
    """
    # 1. Fetch dealers
    dealers_query = db.query(DealerEntity)
    if region and region.strip():
        dealers_query = dealers_query.filter(DealerEntity.region == region.strip())
    if assigned_sale_id:
        dealers_query = dealers_query.filter(DealerEntity.assigned_sale_id == assigned_sale_id)
    if search and search.strip():
        s = f"%{search.strip().lower()}%"
        dealers_query = dealers_query.filter(
            (DealerEntity.name.ilike(s)) | (DealerEntity.code.ilike(s)) | (DealerEntity.phone.ilike(s))
        )
    dealers = dealers_query.all()

    # Cache users for sale staff names
    sales_users = {u.id: (u.full_name or u.username) for u in db.query(UserEntity).all()}

    # 2. Fetch all relevant orders and invoices
    orders = db.query(OrderEntity).filter(
        ~OrderEntity.status.in_(["CANCELLED", "DRAFT", "REJECTED"])
    ).all()

    # Map orders by dealer_id
    dealer_orders_map = {}
    for o in orders:
        if o.dealer_id not in dealer_orders_map:
            dealer_orders_map[o.dealer_id] = []
        dealer_orders_map[o.dealer_id].append(o)

    # Map invoices by dealer_id
    all_invoices = db.query(InvoiceEntity).all()
    dealer_invoices_map = {}
    for inv in all_invoices:
        if inv.dealer_id not in dealer_invoices_map:
            dealer_invoices_map[inv.dealer_id] = []
        dealer_invoices_map[inv.dealer_id].append(inv)

    # 3. Fetch confirmed payment collections (if any)
    dealer_payments_map = {}
    try:
        from app.models.payment_collection import PaymentCollectionEntity
        payment_collections = (
            db.query(PaymentCollectionEntity)
            .filter(PaymentCollectionEntity.status == "CONFIRMED")
            .all()
        )
        for pc in payment_collections:
            if pc.dealer_id:
                amt = float(getattr(pc, "amount", getattr(pc, "amount_collected", 0.0)) or 0.0)
                dealer_payments_map[pc.dealer_id] = (
                    dealer_payments_map.get(pc.dealer_id, 0.0) + amt
                )
    except Exception:
        dealer_payments_map = {}

    now = datetime.now(timezone.utc)
    items: List[DealerDebtAgeingItem] = []

    for d in dealers:
        d_invoices = dealer_invoices_map.get(d.id, [])
        d_orders = dealer_orders_map.get(d.id, [])

        total_receivable = 0.0
        total_paid = 0.0
        remaining_debt = 0.0

        b_0_30 = 0.0
        b_31_60 = 0.0
        b_61_90 = 0.0
        b_over_90 = 0.0
        oldest_days = 0

        if d_invoices:
            # Ưu tiên tính trực tiếp từ dữ liệu hoá đơn đã đồng bộ và đối trừ
            for inv in d_invoices:
                tot = float(inv.total_amount or 0.0)
                paid = float(inv.paid_amount or 0.0)
                rem = max(0.0, float(inv.remaining_debt or (tot - paid)))

                total_receivable += tot
                total_paid += paid
                remaining_debt += rem

                if rem > 0:
                    created_dt = inv.issue_date or inv.created_at
                    if created_dt:
                        if created_dt.tzinfo is None:
                            created_dt = created_dt.replace(tzinfo=timezone.utc)
                        created_date = created_dt.date() if hasattr(created_dt, "date") else created_dt
                        now_date = now.date() if hasattr(now, "date") else now
                        calendar_days = max(0, (now_date - created_date).days)
                        time_days = max(0, (now - created_dt).days)
                        age_days = max(calendar_days, time_days)
                    else:
                        age_days = 0

                    if age_days > oldest_days:
                        oldest_days = age_days

                    if age_days <= 30:
                        b_0_30 += rem
                    elif age_days <= 60:
                        b_31_60 += rem
                    elif age_days <= 90:
                        b_61_90 += rem
                    else:
                        b_over_90 += rem
        else:
            unpaid_orders = []
            for o in d_orders:
                amt = float(o.total_amount or 0.0)
                total_receivable += amt

                if o.status == "PAID":
                    total_paid += amt
                else:
                    created_dt = o.created_at
                    if created_dt:
                        if created_dt.tzinfo is None:
                            created_dt = created_dt.replace(tzinfo=timezone.utc)
                        created_date = created_dt.date() if hasattr(created_dt, "date") else created_dt
                        now_date = now.date() if hasattr(now, "date") else now
                        calendar_days = max(0, (now_date - created_date).days)
                        time_days = max(0, (now - created_dt).days)
                        age_days = max(calendar_days, time_days)
                    else:
                        age_days = 0
                    unpaid_orders.append({
                        "amount": amt,
                        "age_days": age_days,
                        "created_at": created_dt,
                    })

            # Deduct confirmed collections from route collections
            extra_paid = dealer_payments_map.get(d.id, 0.0)
            total_paid += extra_paid

            # Cấn trừ số tiền thu theo tuyến (extra_paid) vào các đơn chưa thanh toán theo FIFO (đơn cũ trước)
            unpaid_orders.sort(key=lambda x: x["created_at"] or datetime.min.replace(tzinfo=timezone.utc))

            rem_pay = extra_paid
            for ord_it in unpaid_orders:
                amt = ord_it["amount"]
                if rem_pay > 0:
                    ded = min(amt, rem_pay)
                    amt -= ded
                    rem_pay -= ded

                if amt > 0.0001:
                    remaining_debt += amt
                    age_days = ord_it["age_days"]
                    if age_days > oldest_days:
                        oldest_days = age_days

                    if age_days <= 30:
                        b_0_30 += amt
                    elif age_days <= 60:
                        b_31_60 += amt
                    elif age_days <= 90:
                        b_61_90 += amt
                    else:
                        b_over_90 += amt

        if remaining_debt <= 0:
            remaining_debt = 0.0
            oldest_days = 0
            b_0_30 = 0.0
            b_31_60 = 0.0
            b_61_90 = 0.0
            b_over_90 = 0.0
            is_overdue = False
        else:
            # Overdue condition:
            # If dealer has debt that reached/exceeded allowed credit terms (max_debt_days)
            # Or has debt past 30 days
            allowed_grace = d.max_debt_days or 30
            is_overdue = (oldest_days >= allowed_grace) or ((b_31_60 + b_61_90 + b_over_90) > 0)

        # Apply overdue_only filter if requested
        if overdue_only and not is_overdue:
            continue

        # Apply ageing_bucket filter if requested
        if ageing_bucket and ageing_bucket not in ["", "all"]:
            if ageing_bucket == "0-30" and b_0_30 <= 0:
                continue
            elif ageing_bucket == "31-60" and b_31_60 <= 0:
                continue
            elif ageing_bucket == "61-90" and b_61_90 <= 0:
                continue
            elif ageing_bucket == "over_90" and b_over_90 <= 0:
                continue
            elif ageing_bucket == "has_debt" and remaining_debt <= 0:
                continue
            elif ageing_bucket == "settled" and remaining_debt > 0:
                continue
            elif ageing_bucket == "overdue" and not is_overdue:
                continue

        sale_name = sales_users.get(d.assigned_sale_id, "Chưa phân công") if d.assigned_sale_id else "Chưa phân công"

        item = DealerDebtAgeingItem(
            dealer_id=d.id,
            dealer_code=d.code,
            dealer_name=d.name,
            region=d.region or "Toàn quốc",
            assigned_sale_id=d.assigned_sale_id,
            assigned_sale_name=sale_name,
            phone=d.phone or "",
            credit_limit=float(d.credit_limit or 0.0),
            max_debt_days=d.max_debt_days or 30,
            overdue_days_allowed=d.overdue_days_allowed or 0,
            total_receivable=total_receivable,
            total_paid=total_paid,
            remaining_debt=remaining_debt,
            oldest_debt_days=oldest_days,
            is_overdue=is_overdue,
            ageing=AgeingBuckets(
                bucket_0_30=b_0_30,
                bucket_31_60=b_31_60,
                bucket_61_90=b_61_90,
                bucket_over_90=b_over_90,
            ),
        )
        items.append(item)

    # Sort: put highest debt and overdue items first
    items.sort(key=lambda x: (not x.is_overdue, -x.remaining_debt, -x.oldest_debt_days))

    # Summary
    summary = DebtLedgerSummary(
        total_dealers=len(items),
        total_receivable=sum(x.total_receivable for x in items),
        total_paid=sum(x.total_paid for x in items),
        total_remaining_debt=sum(x.remaining_debt for x in items),
        total_bucket_0_30=sum(x.ageing.bucket_0_30 for x in items),
        total_bucket_31_60=sum(x.ageing.bucket_31_60 for x in items),
        total_bucket_61_90=sum(x.ageing.bucket_61_90 for x in items),
        total_bucket_over_90=sum(x.ageing.bucket_over_90 for x in items),
    )

    return DebtLedgerResponse(items=items, summary=summary)


def export_debt_ledger_excel(data: DebtLedgerResponse) -> bytes:
    """
    Generates a professional Excel (.xlsx) file for the debt ledger and ageing report.
    """
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "So Cong No & Tuoi No"

    # Styling definitions
    title_font = Font(name="Segoe UI", size=15, bold=True, color="1E3A8A")
    subtitle_font = Font(name="Segoe UI", size=10, italic=True, color="4B5563")
    header_font = Font(name="Segoe UI", size=11, bold=True, color="FFFFFF")
    data_font = Font(name="Segoe UI", size=10)
    summary_font = Font(name="Segoe UI", size=10, bold=True)
    overdue_font = Font(name="Segoe UI", size=10, bold=True, color="DC2626")

    header_fill = PatternFill(start_color="1E40AF", end_color="1E40AF", fill_type="solid")
    bucket_header_fill = PatternFill(start_color="B45309", end_color="B45309", fill_type="solid") # Amber
    summary_fill = PatternFill(start_color="FEF3C7", end_color="FEF3C7", fill_type="solid") # Light yellow

    thin_border = Border(
        left=Side(style="thin", color="D1D5DB"),
        right=Side(style="thin", color="D1D5DB"),
        top=Side(style="thin", color="D1D5DB"),
        bottom=Side(style="thin", color="D1D5DB"),
    )
    summary_border = Border(
        left=Side(style="thin", color="9CA3AF"),
        right=Side(style="thin", color="9CA3AF"),
        top=Side(style="medium", color="4B5563"),
        bottom=Side(style="double", color="111827"),
    )

    # Title block
    ws.merge_cells("A1:K1")
    ws["A1"] = "BÁO CÁO SỔ CÔNG NỢ THEO ĐẠI LÝ VÀ PHÂN TÍCH TUỔI NỢ"
    ws["A1"].font = title_font
    ws["A1"].alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[1].height = 32

    ws.merge_cells("A2:K2")
    ws["A2"] = f"Ngày xuất báo cáo: {datetime.now().strftime('%d/%m/%Y %H:%M')} | Hệ thống Quản trị Bán hàng & Phân phối"
    ws["A2"].font = subtitle_font
    ws["A2"].alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[2].height = 20

    # Headers
    headers = [
        ("STT", 6),
        ("Mã ĐL", 12),
        ("Tên Đại Lý", 30),
        ("Khu Vực", 16),
        ("Phụ Trách", 20),
        ("Tổng Phải Thu (đ)", 18),
        ("Tổng Đã Thu (đ)", 18),
        ("Tổng Còn Nợ (đ)", 18),
        ("0 - 30 Ngày (đ)", 16),
        ("31 - 60 Ngày (đ)", 16),
        ("61 - 90 Ngày (đ)", 16),
        ("> 90 Ngày (đ)", 16),
        ("Trạng Thái Nợ", 16),
    ]

    header_row = 4
    ws.row_dimensions[header_row].height = 28

    for col_idx, (header_text, col_width) in enumerate(headers, start=1):
        cell = ws.cell(row=header_row, column=col_idx)
        cell.value = header_text
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        # Highlighting ageing buckets with distinct color
        if col_idx in [9, 10, 11, 12]:
            cell.fill = bucket_header_fill
        else:
            cell.fill = header_fill
        cell.border = thin_border
        col_letter = get_column_letter(col_idx)
        ws.column_dimensions[col_letter].width = col_width

    # Data Rows
    current_row = 5
    for idx, item in enumerate(data.items, start=1):
        ws.row_dimensions[current_row].height = 22
        status_text = "Quá hạn" if item.is_overdue else "Trong hạn"

        row_values = [
            (idx, Alignment(horizontal="center")),
            (item.dealer_code, Alignment(horizontal="center")),
            (item.dealer_name, Alignment(horizontal="left")),
            (item.region, Alignment(horizontal="center")),
            (item.assigned_sale_name, Alignment(horizontal="left")),
            (item.total_receivable, Alignment(horizontal="right")),
            (item.total_paid, Alignment(horizontal="right")),
            (item.remaining_debt, Alignment(horizontal="right")),
            (item.ageing.bucket_0_30, Alignment(horizontal="right")),
            (item.ageing.bucket_31_60, Alignment(horizontal="right")),
            (item.ageing.bucket_61_90, Alignment(horizontal="right")),
            (item.ageing.bucket_over_90, Alignment(horizontal="right")),
            (status_text, Alignment(horizontal="center")),
        ]

        for col_idx, (val, align) in enumerate(row_values, start=1):
            cell = ws.cell(row=current_row, column=col_idx)
            cell.value = val
            cell.alignment = align
            cell.border = thin_border
            cell.font = data_font

            # Money formatting
            if col_idx in [6, 7, 8, 9, 10, 11, 12]:
                cell.number_format = "#,##0"

            # Color overdue cells
            if col_idx == 13 and item.is_overdue:
                cell.font = overdue_font
            if col_idx in [10, 11, 12] and val > 0:
                cell.font = overdue_font

        current_row += 1

    # Summary Row
    ws.row_dimensions[current_row].height = 26
    ws.merge_cells(start_row=current_row, start_column=1, end_row=current_row, end_column=5)
    summary_label_cell = ws.cell(row=current_row, column=1)
    summary_label_cell.value = f"TỔNG CỘNG ({data.summary.total_dealers} Đại lý)"
    summary_label_cell.font = summary_font
    summary_label_cell.alignment = Alignment(horizontal="center", vertical="center")

    summary_values = [
        (6, data.summary.total_receivable),
        (7, data.summary.total_paid),
        (8, data.summary.total_remaining_debt),
        (9, data.summary.total_bucket_0_30),
        (10, data.summary.total_bucket_31_60),
        (11, data.summary.total_bucket_61_90),
        (12, data.summary.total_bucket_over_90),
    ]

    for col in range(1, 14):
        c = ws.cell(row=current_row, column=col)
        c.fill = summary_fill
        c.border = summary_border

    for col_idx, sum_val in summary_values:
        cell = ws.cell(row=current_row, column=col_idx)
        cell.value = sum_val
        cell.font = summary_font
        cell.number_format = "#,##0"
        cell.alignment = Alignment(horizontal="right", vertical="center")

    output = io.BytesIO()
    wb.save(output)
    return output.getvalue()
